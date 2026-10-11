import os
import docx

def build_design_form1():
    template_docx = r'Standard Templates\2. Design registration\Design registration - Form 1.docx'
    out_dir = r'DESIGN_REGISTRATION'
    os.makedirs(out_dir, exist_ok=True)
    out_docx = os.path.join(out_dir, 'Design_Registration_Form_1_our_project.docx')
    out_pdf = os.path.join(out_dir, 'Design_Registration_Form_1_our_project.pdf')
    
    doc = docx.Document(template_docx)
    
    new_applicants_str = (
        "Class: 14-02- DATA PROCESSING EQUIPMENT AND PERIPHERALS, COMMUNICATION EQUIPMENT. in the name of:\n\n"
        "KAYALVIZHI, Indian, Department of Computer Science and Engineering, Bannari Amman Institute of Technology, Sathyamangalam, Erode – 638401, Tamil Nadu\n\n"
        "VENMUGIL RAJAN S, Indian, Department of Computer Science and Engineering, Bannari Amman Institute of Technology, Sathyamangalam, Erode – 638401, Tamil Nadu\n\n"
        "BUVANESH C, Indian, Department of Computer Science and Engineering, Bannari Amman Institute of Technology, Sathyamangalam, Erode – 638401, Tamil Nadu\n\n"
        "HARISH P, Indian, Department of Computer Science and Engineering, Bannari Amman Institute of Technology, Sathyamangalam, Erode – 638401, Tamil Nadu"
    )
    
    # Table 0: Class and Applicants
    for row in doc.tables[0].rows:
        cell = row.cells[1]
        if 'NISHANTH SHANMUGA SUNDARAM' in cell.text:
            cell.text = "You are requested to register the accompanying in " + new_applicants_str
            
    # Table 1: Title of Article
    for row in doc.tables[1].rows:
        for cell in row.cells:
            if 'MANUAL AGRICULTURAL' in cell.text:
                cell.text = "The design is to be applied to “EDGE-AI SIGN LANGUAGE TRANSLATION APPARATUS”."
                
    # Table 2: Date
    for row in doc.tables[2].rows:
        for cell in row.cells:
            if '12 th day of  March 2026' in cell.text:
                cell.text = cell.text.replace('12 th day of  March 2026', '11th day of October 2026')
                
    doc.save(out_docx)
    print(f"Saved: {out_docx}")
    
    try:
        import win32com.client
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        wdoc = word.Documents.Open(os.path.abspath(out_docx))
        wdoc.SaveAs(os.path.abspath(out_pdf), FileFormat=17) # 17 = wdFormatPDF
        wdoc.Close()
        word.Quit()
        print(f"Successfully converted to PDF: {out_pdf}")
    except Exception as e:
        print(f"Word COM PDF conversion notice: {e}")

if __name__ == '__main__':
    build_design_form1()
