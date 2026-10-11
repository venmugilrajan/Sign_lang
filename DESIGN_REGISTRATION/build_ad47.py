import os
import docx

def build_ad47_direct():
    template_docx = r'Standard Templates\2. Design registration\OTHER DOCUMENTS-AD47 new.docx'
    out_dir = r'DESIGN_REGISTRATION'
    os.makedirs(out_dir, exist_ok=True)
    out_docx = os.path.join(out_dir, 'OTHER_DOCUMENTS_AD47_our_project.docx')
    out_pdf = os.path.join(out_dir, 'OTHER_DOCUMENTS_AD47_our_project.pdf')
    
    doc = docx.Document(template_docx)
    
    # P6: Title of the Article
    for r in doc.paragraphs[6].runs:
        if 'MANUAL AGRICULTURAL' in r.text:
            r.text = r.text.replace('MANUAL AGRICULTURAL MATERIAL COLLECTOR WITH INCLINED BELT CONVEYOR', 'EDGE-AI SIGN LANGUAGE TRANSLATION APPARATUS')
            
    # P8: Brief Description
    for r in doc.paragraphs[8].runs:
        if 'The present design relates to a manual' in r.text:
            r.text = 'The present design relates to a standalone, portable Edge-AI sign language translation apparatus comprising an ergonomic integrated housing, an optical sensor and illumination pod, a front interactive display screen for dynamic 3D sign demonstrations and subtitles, directional acoustic ports, rear thermal exhaust vents, and a stable beveled pedestal base for hands-free counter and desktop operation'
            
    # P10: Statement of Novelty
    for r in doc.paragraphs[10].runs:
        if 'MANUAL AGRICULTURAL' in r.text:
            r.text = 'EDGE-AI SIGN LANGUAGE TRANSLATION APPARATUS'
            
    # P14: Date
    for r in doc.paragraphs[14].runs:
        if '12' in r.text:
            r.text = '11'
        if 'day of March' in r.text:
            r.text = 'day of October'
            
    doc.save(out_docx)
    print(f"Saved: {out_docx}")
    
    # Export to PDF via Word COM
    try:
        import win32com.client
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        wdoc = word.Documents.Open(os.path.abspath(out_docx))
        wdoc.SaveAs(os.path.abspath(out_pdf), FileFormat=17) # 17 = wdFormatPDF
        wdoc.Close()
        word.Quit()
        print(f"Successfully converted to PDF: {out_pdf}")
    except Exception as e:
        print(f"Word COM PDF conversion notice: {e}")

if __name__ == '__main__':
    build_ad47_direct()
