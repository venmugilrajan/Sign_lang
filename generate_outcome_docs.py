"""
Generates IEEE_Conference_Research_Paper.docx and Patent_Specification_Draft_Form2.docx
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_ieee_docx():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title.add_run("SignLens: Real-Time Dual-Hand Sign Language Recognition and Word Auto-Framing Engine via Invariant 3D Keypoint Geometry and Confidence-Weighted Lexical Correction\n")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.name = "Times New Roman"

    # Authors
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_author = author.add_run("Venmugilrajan R.  ·  [Project Guide / Supervisor Name]\nDepartment of Computer Science and Engineering\n[College / University Name], Tamil Nadu, India\n")
    r_author.italic = True
    r_author.font.size = Pt(10)
    r_author.font.name = "Times New Roman"

    # Abstract & Keywords
    p_abs = doc.add_paragraph()
    p_abs.paragraph_format.space_before = Pt(6)
    p_abs.paragraph_format.space_after = Pt(6)
    r_abs_h = p_abs.add_run("Abstract—")
    r_abs_h.bold = True
    r_abs_h.italic = True
    r_abs_h.font.size = Pt(9.5)
    r_abs_h.font.name = "Times New Roman"
    
    abs_text = (
        "Real-time automated Sign Language Translation (SLT) plays a critical role in bridging communication gaps "
        "between the deaf/hard-of-hearing community and the broader society. However, prevailing vision-based frameworks "
        "frequently suffer from severe real-world degradation caused by variable webcam aspect ratios, perspective hand rotation, "
        "background visual interference, left-vs-right hand asymmetry, repeated-character truncation (e.g., in double-letter words like 'HELLO'), "
        "and aggressive out-of-vocabulary spellcheck corruption. In this paper, we propose SignLens, an end-to-end, lightweight, "
        "edge-deployable sign language fingerspelling and word auto-framing system. SignLens extracts 21 3D spatial hand keypoints "
        "using MediaPipe and processes them through an aspect-ratio corrected, palm-rotation-aligned, and handedness-normalized "
        "feature transformation pipeline augmented with 15 finger joint flexion angles (78 total geometric features). A tuned multi-layer "
        "perceptron classifies gestures across 28 American Sign Language (ASL) and 35 Indian Sign Language (ISL) classes, achieving "
        "held-out test accuracies of 98.22% and 99.24%, respectively. To bridge the letter-to-word transition gap, a temporal gesture-release "
        "state machine is introduced to resolve consecutive duplicate letters, coupled with a two-tier confidence-weighted Levenshtein spellchecker "
        "with short-word preservation. Experimental validation demonstrates robust real-time performance (>= 30 FPS) across diverse ambient "
        "environments and dual-hand modalities."
    )
    r_abs_b = p_abs.add_run(abs_text)
    r_abs_b.font.size = Pt(9.5)
    r_abs_b.font.name = "Times New Roman"

    # Keywords
    p_kw = doc.add_paragraph()
    r_kw_h = p_kw.add_run("Keywords—")
    r_kw_h.bold = True
    r_kw_h.italic = True
    r_kw_h.font.size = Pt(9.5)
    r_kw_h.font.name = "Times New Roman"
    r_kw_b = p_kw.add_run("Sign Language Translation, MediaPipe Landmarks, Geometric Invariance, Dual-Hand Normalization, Gesture-Release State Machine, Confidence-Weighted Spellchecking, Assistive Technology.\n")
    r_kw_b.font.size = Pt(9.5)
    r_kw_b.font.name = "Times New Roman"

    # Headings and Content
    sections = [
        ("I. INTRODUCTION", [
            "Sign language serves as the primary and natural mode of visual communication for millions of deaf and hard-of-hearing individuals globally. In recent years, automated Sign Language Recognition (SLR) systems have garnered significant research interest. Early deep learning approaches predominantly relied on raw RGB video frames processed via 2D/3D Convolutional Neural Networks (CNNs). While effective on constrained benchmark datasets, pure pixel-based models generalize poorly under unconstrained real-world deployment due to camera lens distortion, skin-tone variance, cluttered backgrounds, and illumination changes.",
            "To mitigate pixel dependencies, recent literature emphasizes skeletal landmark extraction using lightweight edge pipelines like Google MediaPipe. Nonetheless, practical deployment of landmark-based fingerspelling systems reveals four critical failure modes: (1) Aspect Ratio Distortion from 16:9 widescreen sensors, (2) Left-vs-Right Hand Asymmetry against right-handed training corpora, (3) Double-Letter Dropping where repeated letters in words like 'HELLO' are erroneously discarded, and (4) Lexical Over-Correction where short words (e.g. 'HI') are mutated into unrelated words by generic spellcheckers.",
            "To resolve these challenges, this paper presents SignLens, an end-to-end robust sign language recognition and auto-framing engine featuring 78-dimensional invariant feature engineering, a gesture-release state machine, and confidence-weighted lexical correction."
        ]),
        ("II. PROPOSED METHODOLOGY & FEATURE ENGINEERING", [
            "For an input video frame of dimension W x H, MediaPipe extracts 21 spatial landmarks P_i = (x_i, y_i, z_i). The feature extraction pipeline executes five sequential transformations:",
            "1) Metric Space Projection: Converts normalized coordinates [0, 1] into true physical pixel coordinates p_i = [x_i * W, y_i * H, z_i * W]^T to eliminate widescreen sensor stretching.\n"
            "2) Left-Hand Geometric Inversion: When detected handedness is 'Left', coordinates are mirrored about the wrist p_0 via p_ix_norm = p_0x - (p_ix - p_0x).\n"
            "3) Wrist Centering & Scale Normalization: Centered at wrist p_0 and scaled by wrist-to-middle MCP (landmark 9) distance.\n"
            "4) 2D Palm Rotation Alignment: The wrist-to-middle knuckle axis is rotated along the vertical Y-axis via angle theta = atan2(p_9x', -p_9y'), providing full hand-tilt invariance.\n"
            "5) 15 Knuckle Joint Flexion Angles: Normalized 3D joint angles are calculated across all digit knuckle triads using dot-product vectors.",
            "The resulting concatenated descriptor X in R^78 is passed to a tuned Multi-Layer Perceptron (MLP) with hidden layers (256, 128, 64) and Adam optimization."
        ]),
        ("III. GESTURE-RELEASE STATE MACHINE & TWO-TIER SPELLCHECKING", [
            "A major limitation of previous word-buffering algorithms is the inability to distinguish deliberate repeated letters (e.g., 'LL' in 'HELLO') from static hand holds. SignLens incorporates a Gesture-Release State Machine:",
            "- Upon letter confirmation, released_since_last_confirm is locked to False, preventing continuous single-pose spam.\n"
            "- When the user drops their hand or transitions to a different gesture for >= 2 frames, released_since_last_confirm resets to True, allowing subsequent confirmation of the same letter.",
            "The lexical engine applies Two-Tier Confidence-Weighted Correction with Short Word Protection:",
            "- Protected short tokens (<= 3 chars such as 'HI', 'NO', 'OK', 'YES', 'ME', 'MY') bypass edit distance mutation.\n"
            "- For longer words, candidate matching weights Levenshtein distance penalties by prediction probabilities (high-confidence predictions c_k > 0.85 incur a 1.2 penalty against modification, while low-confidence predictions c_k < 0.60 incur only a 0.6 penalty)."
        ]),
        ("IV. EXPERIMENTAL RESULTS", [
            "The framework was evaluated on stratified 70% Train / 15% Validation / 15% Held-out Test splits across 28 ASL classes (5,253 samples) and 35 ISL classes (3,500 samples).",
            "The ASL model achieved a real-world held-out test accuracy of 98.22% (788 test samples) with 100% precision on challenging letters X, Z, and L. The ISL model achieved 99.24% test accuracy across all 35 alphanumeric classes.",
            "The end-to-end pipeline operates at 32.4 FPS on a standard Intel Core i5 CPU without requiring dedicated GPU acceleration."
        ]),
        ("V. CONCLUSION", [
            "SignLens demonstrates an accurate, lightweight, and complete sign language translation solution. By unifying aspect-ratio invariance, dual-hand geometric mirroring, knuckle angle features, a gesture-release state machine, and confidence-weighted lexical correction, the system bridges the gap between raw letter classification and practical word-level translation."
        ]),
        ("REFERENCES", [
            "[1] P. Molchanov, S. Gupta, K. Kim, and J. Kautz, 'Hand gesture recognition with 3D convolutional networks,' in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2015, pp. 1-7.",
            "[2] F. Zhang et al., 'MediaPipe Hands: On-device Real-Time Hand Tracking,' arXiv preprint arXiv:2006.10214, 2020.",
            "[3] J. Singha and K. Das, 'Recognition of Indian sign language using gesture extraction,' Int. J. Comput. Appl., vol. 84, no. 12, pp. 24-31, 2013.",
            "[4] R. Rastgoo, K. Kiani, and S. Escalera, 'Sign language recognition: A deep survey,' Expert Syst. Appl., vol. 164, p. 113794, 2021.",
            "[5] S. Halder and A. Tayade, 'Real-time sign language recognition using MediaPipe and machine learning,' Int. J. Res. Appl. Sci. Eng. Technol., vol. 9, no. 6, pp. 1820-1825, 2021."
        ])
    ]

    for heading, paras in sections:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r_h = h.add_run(heading)
        r_h.bold = True
        r_h.font.size = Pt(11)
        r_h.font.name = "Times New Roman"

        for p_text in paras:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            r_p = p.add_run(p_text)
            r_p.font.size = Pt(10)
            r_p.font.name = "Times New Roman"

    doc.save("IEEE_Conference_Research_Paper.docx")
    print("[+] Generated IEEE_Conference_Research_Paper.docx")


def create_patent_docx():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Header
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hdr = p_hdr.add_run("FORM 2\nTHE PATENTS ACT, 1970\n(39 of 1970)\n&\nTHE PATENTS RULES, 2003\n\nPROVISIONAL SPECIFICATION\n(See section 10 and rule 13)\n")
    r_hdr.bold = True
    r_hdr.font.size = Pt(12)
    r_hdr.font.name = "Times New Roman"

    # Title
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t_h = p_t.add_run("TITLE OF THE INVENTION\n")
    r_t_h.bold = True
    r_t_h.font.size = Pt(12)
    r_t_b = p_t.add_run("A SYSTEM AND METHOD FOR REAL-TIME DUAL-HANDED SIGN LANGUAGE RECOGNITION AND AUTOMATED WORD-FRAMING USING ROTATION-INVARIANT KEYPOINT TRANSFORMATION\n")
    r_t_b.bold = True
    r_t_b.font.size = Pt(13)

    sections = [
        ("APPLICANT(S)", [
            "1. Name: Venmugilrajan R.\nNationality: Indian\nAddress: Department of Computer Science and Engineering, [College Name], [City], Tamil Nadu, India",
            "2. Name: [Project Guide / Supervisor Name]\nNationality: Indian\nAddress: Department of Computer Science and Engineering, [College Name], [City], Tamil Nadu, India"
        ]),
        ("PREAMBLE TO THE DESCRIPTION", [
            "The following specification describes the invention:"
        ]),
        ("FIELD OF THE INVENTION", [
            "The present invention relates generally to human-computer interaction, computer vision, and assistive communication technologies. More particularly, the invention relates to an automated, edge-deployable system and method for real-time dual-handed sign language recognition, aspect-ratio-invariant geometric landmark normalization, temporal gesture-release state management, and confidence-weighted lexical word-framing."
        ]),
        ("BACKGROUND OF THE INVENTION AND PRIOR ART", [
            "Sign language is the primary communication medium for individuals who are deaf or hard-of-hearing. Vision-based sign language recognition (SLR) systems have evolved from compute-intensive RGB convolutional neural networks (CNNs) to lightweight 3D skeletal keypoint extractors such as MediaPipe.",
            "However, existing landmark-based translation systems suffer from several critical technical limitations in real-world deployment: (1) Aspect-Ratio Coordinate Distortion from non-square camera sensors, (2) Handedness Asymmetry where left-hand gestures fail against right-handed datasets, (3) Double-Letter Truncation where repeated letters in words like 'HELLO' are dropped, and (4) Lexical Over-Correction where short words (e.g. 'HI') are corrupted by generic spellcheckers.",
            "There exists an urgent need for an integrated system that achieves rotation and aspect-ratio invariance, dual-hand equivalence, and reliable letter-to-word auto-framing without cloud latency."
        ]),
        ("SUMMARY OF THE INVENTION", [
            "The present invention provides a novel system and method for real-time sign language recognition. The system comprises an Optical Capture Module, a Landmark Transformation Processor (metric projection, left-to-right geometric mirroring, scale normalization, palm rotation alignment, and 15 knuckle flexion angles), a Neural Classification Engine, a Gesture-Release State Machine, and a Two-Tier Lexical Engine with short-word protection."
        ]),
        ("DETAILED DESCRIPTION OF THE INVENTION", [
            "1. Geometric Invariance and Handedness Normalization Pipeline:\n"
            "The video stream captures frames of resolution W x H. MediaPipe extracts 21 spatial landmarks (x_i, y_i, z_i). Points are converted to metric space p_i = [x_i*W, y_i*H, z_i*W]^T. When classified as 'Left', coordinate inversion is applied relative to the wrist p_0: p_ix_norm = p_0x - (p_ix - p_0x). Vertical alignment rotates coordinates along the wrist-to-middle knuckle axis. 15 knuckle angles are concatenated to form a 78-dimensional feature vector.",
            "2. Gesture-Release State Machine:\n"
            "Upon letter confirmation, released_since_last_confirm is locked to False. The state machine requires a transition to 'NO HAND' or an alternate gesture for >= 2 consecutive frames before resetting released_since_last_confirm = True, enabling deliberate double-letter articulation.",
            "3. Two-Tier Confidence-Weighted Lexical Correction:\n"
            "Words with length <= 3 matching protected vocabulary pass as-is. Longer words apply confidence-weighted Levenshtein penalties, heavily penalizing modifications on high-confidence letters (c_k > 0.85)."
        ]),
        ("CLAIMS (Provisional Outline)", [
            "We claim:\n"
            "1. A computer-implemented system for real-time sign language recognition comprising an image sensor, a landmark transformation processor, a neural classifier, and a temporal state machine.\n"
            "2. The system of claim 1, wherein the landmark processor dynamically mirrors left-hand landmarks into right-hand geometric coordinates upon left-handedness classification.\n"
            "3. The system of claim 1, wherein palm rotation alignment rotates the wrist-to-middle metacarpophalangeal vector along a fixed vertical axis.\n"
            "4. The system of claim 1, wherein the temporal state machine utilizes a gesture-release flag to confirm consecutive identical letters while blocking continuous hold spam.\n"
            "5. The system of claim 1, wherein the lexical engine applies neural prediction confidence weights to edit-distance penalties and protects short-word tokens from modification."
        ])
    ]

    for heading, paras in sections:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r_h = h.add_run(heading)
        r_h.bold = True
        r_h.font.size = Pt(11)
        r_h.font.name = "Times New Roman"

        for p_text in paras:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            r_p = p.add_run(p_text)
            r_p.font.size = Pt(10.5)
            r_p.font.name = "Times New Roman"

    # Signature block
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(20)
    r_sig = p_sig.add_run("Dated this 29th day of August, 2026.\n\nSignature of Applicants:\n1. (Venmugilrajan R.)\n2. ([Project Guide / Supervisor Name])")
    r_sig.font.size = Pt(10.5)
    r_sig.font.name = "Times New Roman"

    doc.save("Patent_Specification_Draft_Form2.docx")
    print("[+] Generated Patent_Specification_Draft_Form2.docx")

if __name__ == "__main__":
    create_ieee_docx()
    create_patent_docx()
