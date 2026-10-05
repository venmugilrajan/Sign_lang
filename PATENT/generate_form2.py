"""
Clone the sample Form 2.docx template and replace ONLY the yellow-highlighted
invention-specific content, keeping all statutory headers, margins, styles, and
formatting identical to the IPR Cell template.
"""
import copy, re
import docx

SRC = 'r:/Projects/Sign Language/PATENT/Form 2.docx'
DST = 'r:/Projects/Sign Language/PATENT/Form 2_our_project.docx'

doc = docx.Document(SRC)

# ── Helper: replace text in a cell while preserving formatting ──
def replace_cell_text(cell, old_frag, new_frag):
    for p in cell.paragraphs:
        for run in p.runs:
            if old_frag in run.text:
                run.text = run.text.replace(old_frag, new_frag)

# ── Helper: replace text inside a paragraph (all runs) ──
def replace_para_text(para, new_text):
    """Replace all run text with new_text, keeping first run's formatting."""
    if not para.runs:
        return
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ''

def set_para_text(para, new_text):
    """Set paragraph text preserving XML structure of runs."""
    full = ''.join(r.text for r in para.runs)
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ''

# ═══════════════════════════════════════════════════════════
# 1. TABLE HEADER (rows 0-4)
# ═══════════════════════════════════════════════════════════
tbl = doc.tables[0]

# Row 1: Title
OLD_TITLE = 'A CLOUD-BASED INTELLIGENT IRRIGATION DECISION SYSTEM WITH UNCERTAINTY-AWARE CONTROL AND UTILITY POWER FALLBACK'
NEW_TITLE = 'A REAL-TIME VISION-BASED DUAL-HAND SIGN LANGUAGE RECOGNITION SYSTEM USING GEOMETRIC LANDMARKS AND MULTI-LAYER PERCEPTRON'
replace_cell_text(tbl.cell(1, 0), OLD_TITLE, NEW_TITLE)

# Row 2: Applicants (minor cleanup - keep BIT address identical)
# Already says BIT, Sathyamangalam - keep as is

# ═══════════════════════════════════════════════════════════
# 2. BODY PARAGRAPHS – replace content paragraph by paragraph
# ═══════════════════════════════════════════════════════════

# Build paragraph replacement map: index -> new text
# Using the same [XX] numbering style as the sample

PARA_MAP = {
# FIELD OF THE INVENTION
11: '[01] The present invention relates generally to computer vision, edge-computed machine learning, and assistive human-computer interaction technologies. More specifically, the present invention relates to an edge-executable, real-time vision-based dual-hand sign language recognition and translation system configured to extract coordinate-normalized 3D geometric hand landmarks and inter-joint flexion angles, perform high-speed neural classification via a lightweight multi-layer perceptron architecture, and output multi-modal textual and acoustic feedback with zero reliance on cloud infrastructure.',

# BACKGROUND OF THE INVENTION
14: '[02] Sign language serves as the primary and indispensable medium of interpersonal communication for millions of hearing-impaired and speech-impaired individuals globally. Both Indian Sign Language (ISL) and American Sign Language (ASL) incorporate rich grammatical structures characterized by intricate manual gestures, fingertip articulatory postures, and complex spatial configurations involving one or both hands.',

16: '[03] Existing automated sign language recognition systems predominantly employ heavy end-to-end deep learning architectures, such as Convolutional Neural Networks (CNNs), 3D Spatio-Temporal Convolutions (3D-CNNs), and Vision Transformers (ViTs). These approaches consume raw two-dimensional video frames (RGB pixels) directly and attempt to simultaneously learn foreground hand shapes alongside uncontrolled background textures, skin-tone variations, ambient illumination shifts, and environmental shadows. Consequently, such deep pixel-level models suffer from immense computational overhead, large parameter storage sizes (frequently exceeding 300 to 500 Megabytes), high thermal dissipation, and severe latency bottlenecks that prohibit interactive execution on commodity edge hardware, smartphones, and low-power microcomputing modules.',

18: '[04] Furthermore, existing vision-based sign language translation systems exhibit high vulnerability to temporal frame leakage and overfitting during training. Because adjacent frames in continuous video streams share high spatial correlation, traditional randomly shuffled validation techniques produce artificially inflated accuracy metrics that rapidly degrade during real-world deployments across independent signers and variable camera distances.',

20: '[05] In addition, dual-hand sign language gestures, such as those present extensively in Indian Sign Language (ISL), present severe structural recognition challenges. Conventional architectures fail to maintain robust spatial ordering between dominant and non-dominant hands, frequently suffering from hand-permutation confusion, missing-hand vector dimension mismatches, and severe jitter during rapid transition states between consecutive sign lemmas.',

22: '[06] Accordingly, there exists a critical technical need for an intelligent, edge-executable sign language recognition system that eliminates raw-pixel background dependencies, achieves high classification accuracy across dual-hand and single-hand gesture alphabets, maintains scale and distance invariance, and delivers real-time inference at 60+ frames per second on consumer-grade computing hardware.',

# OBJECTIVES OF THE PRESENT INVENTION
25: '[07] The primary objective of the present invention is to provide a real-time vision-based dual-hand sign language recognition system that utilizes compact 3D geometric landmarks and inter-joint angular kinematics to achieve high-accuracy translation with ultra-low computational latency.',

26: '[08] Another objective of the present invention is to provide an apparatus and method that normalizes 3D spatial hand coordinates relative to a dynamic wrist origin and palm-vector scale metric, rendering classification completely invariant to user hand size and camera-to-subject distance.',

28: '[09] Another objective of the present invention is to provide a spatial dual-hand formatting mechanism that dynamically orders primary and secondary hand coordinates into a consistent 156-dimensional feature vector, seamlessly handling both single-hand and dual-hand gestures while preventing dimensionality mismatch via standardized zero-padding.',

30: '[10] Another objective of the present invention is to deploy a compact Multi-Layer Perceptron (MLP) neural classifier that delivers 60+ FPS inference directly on consumer central processing units (CPUs) and edge devices without requiring high-power graphics processing units (GPUs) or external cloud infrastructure.',

32: '[011] Another objective of the present invention is to provide a temporal consensus filtering and debounce buffer that suppresses inter-frame transitional flicker and prevents false classifications during hand movement transitions.',

34: '[012] Another objective of the present invention is to provide multi-modal auditory and textual communication feedback, enabling seamless, bidirectional, and barrier-free interaction between hearing-impaired individuals and non-signing persons.',

36: '[013] Another objective of the present invention is to provide a scalable and interoperable system that supports both Indian Sign Language (ISL) with 35 alphanumeric classes and American Sign Language (ASL) with 28 classes using configurable model weights on the same hardware platform.',

# SUMMARY OF THE INVENTION
39: '[014] The present invention provides a real-time vision-based dual-hand sign language recognition system and method configured for low-latency edge deployment. The system comprises an optical image sensing unit (202), a landmark localization and pipeline module (203), a geometric normalization and scaling module (204), a spatial dual-hand alignment and feature formatter (205), a multi-layer perceptron (MLP) inference engine (206), a temporal consistency and debounce filter (207), an interactive visual display subsystem (208), an acoustic synthesizer subsystem (209), and an edge processing unit (210).',

41: '[015] The optical image sensing unit (202) continuously captures a live video stream of a user executing sign language gestures. The landmark localization module (203) processes incoming video frames in real time to isolate one or two hands and extract twenty-one three-dimensional geometric joint coordinates per hand.',

43: '[016] The geometric normalization and scaling module (204) transforms the extracted coordinates by translating the wrist keypoint to a defined origin and scaling all intermediate phalangeal vectors by an instantaneous palm metric, ensuring complete scale and distance invariance.',

45: '[017] The spatial dual-hand alignment and feature formatter (205) calculates fifteen inter-joint flexion angles across kinematic hand joints, arranges detected hands into ordered primary and secondary spatial slots, and compiles a comprehensive 156-dimensional feature representation (or a 78-dimensional representation for single-hand alphabets).',

47: '[018] The multi-layer perceptron (MLP) inference engine (206) evaluates the geometric feature vector using an edge-optimized neural network with rectified linear unit (ReLU) activations, generating normalized class probabilities corresponding to sign language lemmas. The temporal consistency and debounce filter (207) evaluates predictions across a multi-frame temporal buffer to ensure stable symbol convergence.',

49: '[019] Output predictions are simultaneously dispatched to the visual display subsystem (208) as textual transcriptions and to the acoustic synthesizer subsystem (209) as synthesized speech, achieving real-time, bidirectional accessibility on lightweight edge hardware.',

51: '[020] The system delivers recognition accuracy of 98.05 percent on Indian Sign Language (ISL) with 35 alphanumeric classes and 100 percent on American Sign Language (ASL) with 28 classes, while maintaining an operating rate exceeding sixty frames per second on commodity central processing unit hardware.',

# BRIEF DESCRIPTION OF THE ACCOMPANYING DRAWINGS
54: '[021] Figure 1 illustrates an operational flow diagram representing the systematic method for real-time dual-hand sign language landmark recognition and multi-modal translation according to an embodiment of the present invention.',

56: '[022] Figure 2 illustrates a block diagram representing the hardware and functional system architecture of the edge-executable sign language recognition apparatus according to an embodiment of the present invention.',

# DETAILED DESCRIPTION OF THE INVENTION
59: '[023] Referring to Figure 1 and Figure 2, the operational method and technical architecture of the real-time vision-based dual-hand sign language recognition system (201) are described in comprehensive detail. The system operates on edge processing hardware (210) to deliver real-time translation of complex manual gestures into human-interpretable text and synthesized audio.',

61: '[024] The translation process commences at step (101), where the optical image sensing unit (202), comprising a monocular RGB video camera, continuously acquires live digital video frames of a user executing gestures at a nominal capture rate of 30 to 60 frames per second.',

63: '[025] At step (102), the captured frames are passed to the landmark localization module (203). The localization module executes an edge-optimized machine learning pipeline that detects the presence of human hands within the optical field of view and pinpoints twenty-one (21) distinct anatomical keypoint coordinates for each detected hand, representing the wrist, carpals, metacarpals, proximal phalanges, intermediate phalanges, and distal tips of the thumb, index, middle, ring, and pinky fingers.',

65: '[026] At step (103), the localized coordinates are received by the geometric normalization and scaling module (204). The normalization module translates the anatomical wrist landmark (Keypoint 0) to a Cartesian origin (0, 0, 0). Subsequently, a reference palm-span vector is computed between the wrist landmark and the metacarpophalangeal joint of the middle finger (Keypoint 9). All translated landmark coordinates are divided by the Euclidean magnitude of this palm-span vector, thereby producing normalized, scale-invariant spatial coordinates that are completely independent of hand physical dimensions, camera distance, and monocular perspective foreshortening.',

67: '[027] At step (104), the normalized coordinates are processed to extract fifteen (15) inter-joint flexion angles. These angles represent the angular deviation between adjacent skeletal segments across all five digits, providing high-order kinematic descriptors that encode exact finger curl, extension, abduction, and adduction states.',

69: '[028] At step (105), the spatial dual-hand alignment and feature formatter (205) arranges the extracted features into a canonical input vector. For dual-hand Indian Sign Language (ISL), the formatter organizes landmarks into primary (dominant) and secondary (supporting) slots based on horizontal spatial positioning. When only a single hand is present in the field of view, the secondary hand slot is populated with deterministic zero-padding, producing a strictly consistent 156-dimensional feature vector. For single-hand American Sign Language (ASL) operations, a standardized 78-dimensional feature vector is generated.',

71: '[029] At step (106), the formatted geometric feature vector is supplied to the Multi-Layer Perceptron (MLP) inference engine (206). The MLP inference engine incorporates fully connected layers structured with hidden dimensions of 512, 256, and 128 nodes for dual-hand recognition (or 256, 128, and 64 nodes for single-hand recognition), equipped with Rectified Linear Unit (ReLU) activation functions. The final layer computes unnormalized class scores corresponding to the predefined target alphabet (35 classes for ISL alphanumeric signs, or 28 classes for ASL signs).',

73: '[030] At step (107), the class scores are converted into normalized posterior probabilities via a Softmax transformation. A threshold verification sub-routine assesses whether the highest class probability satisfies a predetermined certainty threshold of at least seventy percent. In the event the confidence score falls below the certainty threshold, the frame is classified as an intermediate or transitional gesture and discarded, preventing spurious symbol generation.',

75: '[031] At step (108), accepted classification candidates are passed through a temporal consistency and debounce filter (207). The filter maintains a sliding temporal window of N consecutive video frames. A final sign consensus is established only when a specific sign classification achieves majority agreement across the temporal window, effectively eliminating inter-gesture transitional flicker and high-frequency sensor noise.',

77: '[032] At step (109), the verified sign output is dispatched to the multi-modal feedback subsystems. The visual display subsystem (208) renders the predicted character or phrase on an interactive graphical display, while the acoustic synthesizer subsystem (209) generates real-time synthesized speech phonemes via a text-to-speech engine, enabling seamless, barrier-free communication between the signer and non-signing observers.',
}

# Paragraphs 79-83 in sample are extra detail paragraphs; 
# We'll blank them out (they become empty, same as sample's empty padding paras)
for idx in [79, 81, 83]:
    PARA_MAP[idx] = ''

# Paragraphs 85-111 (Heading 2 style - system architecture detail paras in sample)
# Replace with our system architecture description
ARCH_PARAS = {
85: '[033] The real-time sign language recognition system (201) comprises an optical image sensing unit (202), a landmark localization and pipeline module (203), a geometric normalization and scaling module (204), a spatial dual-hand alignment and feature formatter (205), a multi-layer perceptron (MLP) inference engine (206), a temporal consistency and debounce filter (207), an interactive visual display subsystem (208), an acoustic synthesizer subsystem (209), and an edge processing unit (210).',

87: '[034] The optical image sensing unit (202) serves as the primary visual data acquisition component. The unit comprises a monocular RGB video camera configured to capture digital video frames at a rate of 30 to 60 frames per second. The unit captures the hand gestures of a user performing sign language within its optical field of view.',

89: '[035] The captured video frames are received by the landmark localization and pipeline module (203). The localization module executes a machine learning pipeline to detect one or more human hands within the video frame and extract twenty-one (21) three-dimensional skeletal joint coordinates per detected hand.',

91: '[036] The geometric normalization and scaling module (204) receives the localized joint coordinates. The module translates the wrist landmark (Keypoint 0) to a Cartesian origin and computes a reference palm-span vector between the wrist and the metacarpophalangeal joint of the middle finger (Keypoint 9). All landmark coordinates are divided by the Euclidean magnitude of this palm-span vector, producing scale-invariant coordinates.',

93: '[037] The spatial dual-hand alignment and feature formatter (205) calculates fifteen inter-joint flexion angles representing angular deviations between adjacent skeletal segments across all five digits. For dual-hand gestures, the formatter arranges hand landmarks into spatially ordered primary and secondary slots based on horizontal position, producing a 156-dimensional feature vector. For single-hand gestures, the secondary slot is zero-padded.',

95: '[038] The multi-layer perceptron (MLP) inference engine (206) receives the formatted geometric feature vector and evaluates it through fully connected layers with hidden dimensions of 512, 256, and 128 nodes for dual-hand classification (or 256, 128, and 64 nodes for single-hand classification), equipped with Rectified Linear Unit (ReLU) activation functions.',

97: '[039] The temporal consistency and debounce filter (207) evaluates prediction outputs across a multi-frame sliding temporal buffer. A final sign consensus is established only when a specific classification achieves majority agreement across the temporal window, eliminating transitional flicker.',

99: '[040] The interactive visual display subsystem (208) renders the validated sign classification as textual subtitles on a graphical display interface. The acoustic synthesizer subsystem (209) converts the validated classification output into synthesized speech through a text-to-speech engine.',

101: '[041] The edge processing unit (210) comprises a host controller hardware platform including a central processing unit, system memory, and associated peripherals. The entire inference pipeline executes locally on the edge processing unit without cloud or server dependency, achieving an operating rate exceeding sixty frames per second.',
}

PARA_MAP.update(ARCH_PARAS)

# Blank out remaining architecture paras that sample had (103-111)
for idx in [103, 105, 107, 109, 111]:
    PARA_MAP[idx] = ''

# ═══════════════════════════════════════════════════════════
# 3. CLAIMS
# ═══════════════════════════════════════════════════════════
CLAIMS_MAP = {
134: '1. A real-time vision-based dual-hand sign language recognition system (201) comprising:',
135: 'an optical image sensing unit (202) configured to capture digital video frames of a user performing sign language gestures;',
136: 'a landmark localization module (203) operatively coupled to the optical sensing unit (202) and configured to detect one or more hands and extract twenty-one 3D skeletal joint coordinates per hand; a geometric normalization and scaling module (204) configured to translate wrist joint coordinates to a Cartesian origin and scale all hand joint coordinates by an instantaneous palm dimension vector to achieve scale and distance invariance;',
137: 'a spatial dual-hand alignment and feature formatter (205) configured to extract inter-joint flexion angles and compile a standardized multi-dimensional geometric feature vector;',
138: 'a multi-layer perceptron (MLP) inference engine (206) executed on an edge processing unit (210) to classify the geometric feature vector into sign language alphanumeric classes;',
139: 'a temporal consistency filter (207) configured to debounce and validate classification outputs across consecutive video frames; and',
140: 'at least one output feedback subsystem (208, 209) configured to output the validated sign language classification as textual display or synthesized speech;',
141: '',  # remove extra sub-clause

143: 'wherein outputs of the multi-layer perceptron inference engine (206) are validated through the temporal consistency filter (207) and dispatched simultaneously to the visual display subsystem (208) and acoustic synthesizer subsystem (209) for real-time multi-modal accessibility.',

145: '2. A method for real-time vision-based sign language recognition comprising the steps of:',
146: 'capturing a continuous digital video stream (101) of manual gestures via an optical camera;',
147: 'localizing twenty-one 3D joint landmark coordinates (102) for each detected hand within the video stream;',
148: 'normalizing the localized 3D joint coordinates (103) by translating an anatomical wrist landmark to a coordinate origin and scaling by a reference palm metric;',
149: 'extracting fifteen inter-joint flexion angles (104) across skeletal phalanges of each detected hand;',
150: 'formatting extracted coordinates and angles into an ordered dual-hand feature vector (105) with standardized padding for absent hands;',
151: 'classifying the formatted feature vector using a multi-layer perceptron neural network (106) executed on an edge computing processor;',
152: 'evaluating prediction confidence against a predetermined certainty threshold (107);',
153: 'applying temporal majority consensus filtering (108) across a sliding window of consecutive video frames; and',
154: 'generating multi-modal translated textual and acoustic outputs (109) corresponding to the recognized sign.',

156: '3. The system as claimed in claim 1, wherein the geometric normalization and scaling module (204) determines the palm dimension vector as the Euclidean distance between a wrist keypoint (Keypoint 0) and a middle-finger metacarpophalangeal keypoint (Keypoint 9), thereby providing invariance to user hand size and camera distance.',

158: '4. The system as claimed in claim 1, wherein the spatial dual-hand alignment and feature formatter (205) arranges hand landmarks into spatially ordered primary and secondary slots, generating a 156-dimensional feature vector for dual-hand gestures and populating secondary slots with zero-padded values for single-hand gestures.',

160: '5. The system as claimed in claim 1, wherein the multi-layer perceptron inference engine (206) comprises fully connected feed-forward layers with Rectified Linear Unit (ReLU) activations having node dimensions of 512, 256, and 128 for dual-hand classification, and 256, 128, and 64 for single-hand classification.',

162: '6. The system as claimed in claim 1, wherein the temporal consistency filter (207) maintains a sliding buffer of consecutive frames and suppresses output generation during inter-gesture dynamic transitions until consensus agreement is established.',

164: '7. The system as claimed in claim 1, wherein the multi-modal output feedback subsystems comprise an interactive graphical display subsystem (208) for rendering real-time transcribed text and an acoustic synthesizer subsystem (209) executing text-to-speech auditory synthesis.',

166: '8. The system as claimed in claim 1, wherein the multi-layer perceptron inference engine (206) executes at an operating rate exceeding sixty frames per second directly on commodity central processing unit hardware without external cloud or server dependency.',

168: '9. The method as claimed in claim 2, wherein the step of extracting inter-joint flexion angles (104) computes the angular articulation between adjacent bone segments for thumb, index, middle, ring, and pinky fingers.',

170: '10. The method as claimed in claim 2, wherein classification outputs failing to satisfy the predetermined certainty threshold of at least seventy percent posterior probability are discarded prior to temporal majority consensus filtering.',
}

PARA_MAP.update(CLAIMS_MAP)

# ═══════════════════════════════════════════════════════════
# 4. DATE & SIGNATURE + ABSTRACT
# ═══════════════════════════════════════════════════════════
PARA_MAP[174] = 'Dated this 2nd day of October, 2026                                                       Vinoth Kumar J'

PARA_MAP[179] = NEW_TITLE

PARA_MAP[182] = 'A real-time vision-based dual-hand sign language recognition system (201) and method are disclosed for high-speed edge translation. The system comprises an optical image sensing unit (202), a landmark localization module (203) extracting twenty-one 3D skeletal joint coordinates per hand, a geometric normalization and scaling module (204) providing wrist-centered and palm-dimension invariant coordinates, a spatial dual-hand alignment formatter (205) deriving inter-joint flexion angles into a standardized feature vector, a multi-layer perceptron (MLP) inference engine (206) operating on an edge processing unit (210), a temporal consistency filter (207), an interactive visual display subsystem (208), and an acoustic synthesizer subsystem (209). The method captures video frames (101), pinpoints joint landmarks (102), normalizes coordinates (103), calculates flexion angles (104), compiles an ordered dual-hand feature vector (105), evaluates MLP predictions (106) against a certainty threshold (107), applies temporal debounce filtering (108), and generates multi-modal textual and synthesized speech outputs (109). The invention delivers 60+ FPS translation with 98.05 percent ISL accuracy completely on edge devices without cloud dependency.'

# ═══════════════════════════════════════════════════════════
# APPLY ALL REPLACEMENTS
# ═══════════════════════════════════════════════════════════
for idx, new_text in PARA_MAP.items():
    if idx < len(doc.paragraphs):
        set_para_text(doc.paragraphs[idx], new_text)

doc.save(DST)
print(f'Form 2_our_project.docx created at: {DST}')
print(f'Total paragraphs: {len(doc.paragraphs)}')
